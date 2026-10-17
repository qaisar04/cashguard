"""Runs the public job API over the labelled dataset and reports detection / localization / classification metrics.

Example:
  python scripts/evaluate.py --labels ~/Downloads/labels_typed_v2.xlsx --videos-dir ~/Downloads/open/videos \
      --base-url http://localhost:8080 --api-key change-me --out eval.csv
"""
from __future__ import annotations

import argparse
import csv
import time
from dataclasses import dataclass
from pathlib import Path

import httpx
from openpyxl import load_workbook

API_KEY_HEADER = "X-API-Key"
STATUS_PROCESSING = "processing"
STATUS_DONE = "done"
VERDICT_VIOLATION = "violation"
IMPLEMENTED_TYPES = {
    "workstation_unlocked",
    "valuables_left_unattended",
    "safe_left_open",
    "cash_room_door_open",
    "unauthorized_entry",
    "cash_exposed_to_visitor",
}
COLUMNS = {"video": "video_id", "violation": "violation", "timecode": "timecode_sec", "type": "type_id", "certainty": "certainty"}
TOLERANCES_SEC = (2.0, 5.0)
CYRILLIC_KHA = "Х"
LATIN_X = "X"
VIDEO_SUFFIXES = (".mov", ".mp4")


@dataclass
class Label:
    video_id: str
    violation: bool
    timecode: float | None
    type_id: str | None
    certainty: str | None


@dataclass
class Result:
    label: Label
    status: str
    verdict: str | None
    events: list[dict]
    elapsed_sec: float
    error: str | None = None

    @property
    def predicted_violation(self) -> bool:
        return self.verdict == VERDICT_VIOLATION

    @property
    def timecode_error(self) -> float | None:
        if self.label.timecode is None or not self.events:
            return None
        return min(abs(e["timecode_sec"] - self.label.timecode) for e in self.events)

    @property
    def type_hit(self) -> bool:
        return any(e["type_id"] == self.label.type_id for e in self.events)


def read_labels(path: Path, sheet: int) -> list[Label]:
    ws = load_workbook(path, read_only=True, data_only=True).worksheets[sheet]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    idx = {key: header.index(name) for key, name in COLUMNS.items() if name in header}
    labels: list[Label] = []
    for row in rows:
        video = row[idx["video"]]
        if not video or not str(video).strip().upper().replace(CYRILLIC_KHA, LATIN_X).startswith(LATIN_X + "-"):
            continue
        timecode = row[idx["timecode"]] if "timecode" in idx else None
        labels.append(
            Label(
                video_id=str(video).strip().replace(CYRILLIC_KHA, LATIN_X),
                violation=bool(int(row[idx["violation"]] or 0)),
                timecode=float(timecode) if timecode not in (None, "") else None,
                type_id=(str(row[idx["type"]]).strip() or None) if "type" in idx and row[idx["type"]] else None,
                certainty=(str(row[idx["certainty"]]).strip() or None) if "certainty" in idx and row[idx["certainty"]] else None,
            )
        )
    return labels


def find_video(videos_dir: Path, video_id: str) -> Path | None:
    wanted = video_id.upper().replace(CYRILLIC_KHA, LATIN_X)
    for path in videos_dir.iterdir():
        if path.suffix.lower() in VIDEO_SUFFIXES and path.stem.upper().replace(CYRILLIC_KHA, LATIN_X) == wanted:
            return path
    return None


def run_one(client: httpx.Client, base_url: str, video: Path, timeout_sec: float, poll_sec: float) -> tuple[str, dict, float]:
    started = time.time()
    with video.open("rb") as f:
        res = client.post(f"{base_url}/api/v1/analyze", files={"file": (video.name, f, "video/quicktime")})
    res.raise_for_status()
    job_id = res.json()["job_id"]
    while time.time() - started < timeout_sec:
        body = client.get(f"{base_url}/api/v1/jobs/{job_id}").json()
        if body["status"] != STATUS_PROCESSING:
            return body["status"], body, time.time() - started
        time.sleep(poll_sec)
    return "timeout", {}, time.time() - started


VISITOR_TYPES = ("unauthorized_entry", "cash_exposed_to_visitor")


def report_clean_false_positives(results: list[Result]) -> None:
    """Clean clips must stay clean; the glass wall reflects the cashier, so visitor-based types get their own line."""
    clean = [r for r in results if not r.label.violation]
    by_type: dict[str, list[str]] = {}
    for r in clean:
        for e in r.events:
            by_type.setdefault(e["type_id"], []).append(r.label.video_id)
    print(f"\n== false positives on {len(clean)} clean clips")
    if not by_type:
        print("   none")
    for type_id, videos in sorted(by_type.items()):
        print(f"   {type_id}: {len(videos)} ({', '.join(videos)})")
    visitor_fp = [v for t in VISITOR_TYPES for v in by_type.get(t, [])]
    print(f"   visitor-type check ({' / '.join(VISITOR_TYPES)}): {'OK' if not visitor_fp else 'FAIL ' + ', '.join(visitor_fp)}")


def summarize(results: list[Result], subset_name: str) -> None:
    tp = sum(1 for r in results if r.label.violation and r.predicted_violation)
    fp = sum(1 for r in results if not r.label.violation and r.predicted_violation)
    fn = sum(1 for r in results if r.label.violation and not r.predicted_violation)
    tn = sum(1 for r in results if not r.label.violation and not r.predicted_violation)
    precision = tp / (tp + fp) if tp + fp else 0.0
    recall = tp / (tp + fn) if tp + fn else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    print(f"\n== {subset_name}: n={len(results)}  TP={tp} FP={fp} FN={fn} TN={tn}")
    print(f"   verdict precision={precision:.2f} recall={recall:.2f} f1={f1:.2f}")
    errors = [r.timecode_error for r in results if r.label.violation and r.timecode_error is not None]
    if errors:
        mae = sum(errors) / len(errors)
        hits = {tol: sum(1 for e in errors if e <= tol) / len(errors) for tol in TOLERANCES_SEC}
        print(f"   timecode MAE={mae:.1f}s  " + "  ".join(f"hit@{tol:.0f}s={v:.2f}" for tol, v in hits.items()))
    typed = [r for r in results if r.label.violation and r.label.type_id]
    if typed:
        acc = sum(1 for r in typed if r.type_hit) / len(typed)
        print(f"   type accuracy (label type among predicted)={acc:.2f} over {len(typed)}")
    failed = [r for r in results if r.status != STATUS_DONE]
    slowest = max((r.elapsed_sec for r in results), default=0.0)
    print(f"   failed/timeout={len(failed)}  slowest={slowest:.0f}s")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--labels", type=Path, required=True)
    parser.add_argument("--videos-dir", type=Path, required=True)
    parser.add_argument("--base-url", default="http://localhost:8080")
    parser.add_argument("--api-key", required=True)
    parser.add_argument("--sheet", type=int, default=0)
    parser.add_argument("--only", nargs="*", default=None, help="video ids to run, e.g. X-11 X-22")
    parser.add_argument("--timeout-sec", type=float, default=330)
    parser.add_argument("--poll-sec", type=float, default=3)
    parser.add_argument("--out", type=Path, default=Path("eval.csv"))
    args = parser.parse_args()

    labels = read_labels(args.labels, args.sheet)
    if args.only:
        wanted = {v.upper().replace(CYRILLIC_KHA, LATIN_X) for v in args.only}
        labels = [l for l in labels if l.video_id.upper() in wanted]
    results: list[Result] = []
    with httpx.Client(headers={API_KEY_HEADER: args.api_key}, timeout=60) as client:
        for label in labels:
            video = find_video(args.videos_dir, label.video_id)
            if video is None:
                print(f"{label.video_id}: video not found, skipped")
                continue
            try:
                status, body, elapsed = run_one(client, args.base_url, video, args.timeout_sec, args.poll_sec)
                result = Result(label, status, body.get("verdict"), body.get("violations") or [], elapsed, body.get("error"))
            except httpx.HTTPError as e:
                result = Result(label, "error", None, [], 0.0, str(e))
            results.append(result)
            predicted = ", ".join(f"{e['type_id']}@{e['timecode_sec']}" for e in result.events) or "-"
            print(f"{label.video_id}: gt={'V' if label.violation else '-'}{'@' + str(label.timecode) if label.timecode else ''} "
                  f"{label.type_id or ''} | pred={result.verdict or result.status} {predicted} | {result.elapsed_sec:.0f}s")

    summarize(results, "all videos")
    report_clean_false_positives(results)
    implemented = [r for r in results if not r.label.violation or r.label.type_id in IMPLEMENTED_TYPES]
    summarize(implemented, "implemented types + clean videos")
    for certainty in sorted({r.label.certainty for r in results if r.label.certainty}):
        summarize([r for r in results if r.label.certainty == certainty], f"certainty={certainty}")

    with args.out.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["video_id", "gt_violation", "gt_timecode", "gt_type", "certainty", "status", "verdict", "pred_events", "timecode_error", "elapsed_sec", "error"])
        for r in results:
            writer.writerow([r.label.video_id, int(r.label.violation), r.label.timecode, r.label.type_id, r.label.certainty, r.status, r.verdict,
                             "; ".join(f"{e['type_id']}@{e['timecode_sec']}" for e in r.events), r.timecode_error, round(r.elapsed_sec, 1), r.error])
    print(f"\nper-video results -> {args.out}")


if __name__ == "__main__":
    main()
